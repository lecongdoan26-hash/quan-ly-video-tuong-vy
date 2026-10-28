import os
import shutil
import uuid
import socket
from datetime import datetime
from typing import Optional
from io import BytesIO

from fastapi import FastAPI, Request, Form, UploadFile, File, Query, HTTPException, Header, status
from fastapi.responses import HTMLResponse, FileResponse, StreamingResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from fastapi.middleware.cors import CORSMiddleware
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import database

# Khởi tạo thư mục
import sys
if getattr(sys, "frozen", False):
    BASE_DIR = os.path.dirname(sys.executable)
    BUNDLE_DIR = getattr(sys, "_MEIPASS", BASE_DIR)
else:
    BASE_DIR = os.path.dirname(os.path.abspath(__file__))
    BUNDLE_DIR = BASE_DIR
UPLOAD_DIR = os.path.join(BASE_DIR, "uploads")
STATIC_DIR = os.path.join(BUNDLE_DIR, "static") if os.path.exists(os.path.join(BUNDLE_DIR, "static")) else os.path.join(BASE_DIR, "static")
TEMPLATES_DIR = os.path.join(BUNDLE_DIR, "templates") if os.path.exists(os.path.join(BUNDLE_DIR, "templates")) else os.path.join(BASE_DIR, "templates")

os.makedirs(UPLOAD_DIR, exist_ok=True)
os.makedirs(STATIC_DIR, exist_ok=True)
os.makedirs(os.path.join(STATIC_DIR, "css"), exist_ok=True)
os.makedirs(os.path.join(STATIC_DIR, "js"), exist_ok=True)
os.makedirs(TEMPLATES_DIR, exist_ok=True)

# Khởi tạo database
database.init_db()

app = FastAPI(title="Hệ Thống Kho & Quản Lý Video - Tường Vy Admin")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")
templates = Jinja2Templates(directory=TEMPLATES_DIR)

def get_local_ip():
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(("8.8.8.8", 80))
        ip = s.getsockname()[0]
        s.close()
        return ip
    except Exception:
        return "127.0.0.1"

@app.get("/", response_class=HTMLResponse)
async def home(request: Request):
    local_ip = get_local_ip()
    return templates.TemplateResponse(
        request=request,
        name="index.html",
        context={
            "local_ip": local_ip,
            "today": datetime.now().strftime("%Y-%m-%d")
        }
    )

@app.get("/api/network-info")
async def get_network_info():
    ip = get_local_ip()
    return {
        "local_ip": ip,
        "url_local": "http://localhost:8000",
        "url_lan": f"http://{ip}:8000"
    }

# ================= AUTH & USER MANAGEMENT =================

# 1. Danh sách người dùng hiển thị trên màn hình đăng nhập (Không để lộ PIN)
@app.get("/api/auth/users-list")
async def get_public_users_list():
    return database.get_users_public()

# 2. Xác thực đăng nhập bằng Mã PIN riêng biệt của từng người
@app.post("/api/auth/login")
async def login(
    user_id: int = Form(...),
    pin: str = Form(...)
):
    user = database.verify_user_login(user_id, pin)
    if not user:
        raise HTTPException(
            status_code=401, 
            detail="Mã PIN không chính xác! Vui lòng liên hệ Tường Vy (Quản trị viên) để lấy mã PIN."
        )

    return {
        "status": "success",
        "user": user,
        "message": f"Chào {user['name']}! Đăng nhập thành công."
    }

# 3. Lấy danh sách thành viên kèm Mã PIN (CHỈ DÀNH CHO TƯỜNG VY ADMIN)
@app.get("/api/admin/users")
async def get_admin_users(
    current_user_role: str = Header("editor")
):
    if current_user_role != "admin":
        raise HTTPException(status_code=403, detail="Chỉ Tường Vy (Quản trị viên) mới có quyền xem mục này!")
    return database.get_users_for_admin()

# 4. Tường Vy thêm nhân sự / Sếp mới kèm mã PIN riêng
@app.post("/api/admin/users")
async def add_new_user(
    name: str = Form(...),
    role: str = Form("editor"),
    pin: str = Form("1234"),
    current_user_role: str = Header("editor")
):
    if current_user_role != "admin":
        raise HTTPException(status_code=403, detail="Chỉ Tường Vy mới có quyền thêm nhân sự!")

    success = database.add_user(name, role, pin)
    if not success:
        raise HTTPException(status_code=400, detail="Tên người dùng này đã tồn tại trong hệ thống!")
    return {"status": "success", "message": f"Đã thêm {name} với mã PIN riêng thành công!"}

# 5. Tường Vy đổi mã PIN cho bất kỳ người nào
@app.post("/api/admin/users/{user_id}/update-pin")
async def update_pin(
    user_id: int,
    new_pin: str = Form(...),
    current_user_role: str = Header("editor")
):
    if current_user_role != "admin":
        raise HTTPException(status_code=403, detail="Chỉ Tường Vy mới có quyền đổi mã PIN của người khác!")

    if not new_pin.strip():
        raise HTTPException(status_code=400, detail="Mã PIN không được để trống!")

    database.update_user_pin(user_id, new_pin.strip())
    return {"status": "success", "message": "Đã cập nhật mã PIN mới thành công!"}

# 6. Tường Vy đổi vai trò (Sếp / Nhân viên)
@app.post("/api/admin/users/{user_id}/update-role")
async def update_role(
    user_id: int,
    new_role: str = Form(...),
    current_user_role: str = Header("editor")
):
    if current_user_role != "admin":
        raise HTTPException(status_code=403, detail="Chỉ Tường Vy mới có quyền đổi quyền của người dùng!")

    success = database.update_user_role(user_id, new_role)
    if not success:
        raise HTTPException(status_code=400, detail="Không thể đổi vai trò của Quản trị viên Tường Vy!")
    return {"status": "success", "message": "Đã đổi vai trò thành công!"}

# 7. Tường Vy xóa nhân sự
@app.delete("/api/admin/users/{user_id}")
async def delete_user(
    user_id: int,
    current_user_role: str = Header("editor")
):
    if current_user_role != "admin":
        raise HTTPException(status_code=403, detail="Chỉ Tường Vy mới có quyền xóa người dùng!")

    success = database.delete_user(user_id)
    if not success:
        raise HTTPException(status_code=400, detail="Không thể xóa tài khoản Quản trị viên Tường Vy!")
    return {"status": "success", "message": "Đã xóa người dùng"}

# ================= VIDEO CRUD WITH PERMISSIONS =================
@app.get("/api/stats")
async def get_stats():
    return database.get_stats()

@app.get("/api/videos")
async def list_videos(
    tab: str = Query("archive"),
    search: Optional[str] = None,
    creator: Optional[str] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None
):
    videos = database.get_all_videos(
        tab=tab,
        search=search,
        creator=creator,
        date_from=date_from,
        date_to=date_to
    )
    return videos

@app.get("/api/videos/{video_id}")
async def get_video(video_id: int):
    video = database.get_video_by_id(video_id)
    if not video:
        raise HTTPException(status_code=404, detail="Không tìm thấy video")
    return video

@app.post("/api/videos")
async def create_video_item(
    title: str = Form(...),
    description: str = Form(""),
    creator: str = Form(...),
    posted_tiktok: int = Form(0),
    posted_facebook: int = Form(0),
    posted_youtube: int = Form(0),
    created_date: str = Form(...),
    published_date: str = Form(""),
    video_file: Optional[UploadFile] = File(None),
    current_user_role: str = Header("editor")
):
    # Sếp chỉ có quyền xem, không được thêm video
    if current_user_role == "boss":
        raise HTTPException(status_code=403, detail="Sếp chỉ có quyền xem video, không được thêm video mới!")

    saved_filename = ""
    original_filename = ""
    file_size = 0

    if video_file and video_file.filename:
        original_filename = video_file.filename
        ext = os.path.splitext(original_filename)[1].lower() or ".mp4"
        saved_filename = f"vid_{datetime.now().strftime('%Y%m%d_%H%M%S')}_{uuid.uuid4().hex[:6]}{ext}"
        target_path = os.path.join(UPLOAD_DIR, saved_filename)
        
        with open(target_path, "wb") as buffer:
            while chunk := await video_file.read(1024 * 1024):
                buffer.write(chunk)
                file_size += len(chunk)

    video_data = {
        "title": title,
        "description": description,
        "creator": creator,
        "posted_tiktok": posted_tiktok,
        "posted_facebook": posted_facebook,
        "posted_youtube": posted_youtube,
        "created_date": created_date or datetime.now().strftime("%Y-%m-%d"),
        "published_date": published_date,
        "video_filename": saved_filename,
        "original_filename": original_filename,
        "file_size": file_size
    }

    new_id = database.create_video(video_data)
    return {"status": "success", "id": new_id, "message": "Thêm video thành công"}

@app.post("/api/videos/{video_id}/update")
async def update_video_item(
    video_id: int,
    title: str = Form(...),
    description: str = Form(""),
    creator: str = Form(...),
    posted_tiktok: int = Form(0),
    posted_facebook: int = Form(0),
    posted_youtube: int = Form(0),
    created_date: str = Form(...),
    published_date: str = Form(""),
    video_file: Optional[UploadFile] = File(None),
    current_user_role: str = Header("editor")
):
    # Sếp chỉ có quyền xem, không được sửa video
    if current_user_role == "boss":
        raise HTTPException(status_code=403, detail="Sếp chỉ có quyền xem video, không được chỉnh sửa video!")

    existing = database.get_video_by_id(video_id)
    if not existing:
        raise HTTPException(status_code=404, detail="Không tìm thấy video")

    update_data = {
        "title": title,
        "description": description,
        "creator": creator,
        "posted_tiktok": posted_tiktok,
        "posted_facebook": posted_facebook,
        "posted_youtube": posted_youtube,
        "created_date": created_date,
        "published_date": published_date
    }

    if video_file and video_file.filename:
        old_file = existing.get("video_filename")
        if old_file:
            old_path = os.path.join(UPLOAD_DIR, old_file)
            if os.path.exists(old_path):
                try:
                    os.remove(old_path)
                except Exception:
                    pass

        original_filename = video_file.filename
        ext = os.path.splitext(original_filename)[1].lower() or ".mp4"
        saved_filename = f"vid_{datetime.now().strftime('%Y%m%d_%H%M%S')}_{uuid.uuid4().hex[:6]}{ext}"
        target_path = os.path.join(UPLOAD_DIR, saved_filename)
        
        file_size = 0
        with open(target_path, "wb") as buffer:
            while chunk := await video_file.read(1024 * 1024):
                buffer.write(chunk)
                file_size += len(chunk)

        update_data["video_filename"] = saved_filename
        update_data["original_filename"] = original_filename
        update_data["file_size"] = file_size

    database.update_video(video_id, update_data)
    return {"status": "success", "message": "Cập nhật video thành công"}

@app.post("/api/videos/{video_id}/toggle-channel")
async def toggle_channel(
    video_id: int,
    channel: str = Form(...),
    current_user_role: str = Header("editor")
):
    # Sếp không được đổi kênh đăng
    if current_user_role == "boss":
        raise HTTPException(status_code=403, detail="Sếp chỉ có quyền xem video, không được thay đổi trạng thái kênh!")

    res = database.toggle_video_channel(video_id, channel)
    if not res:
        raise HTTPException(status_code=404, detail="Không tìm thấy video")
    return {"status": "success", "data": res}

@app.delete("/api/videos/{video_id}")
async def delete_video_item(
    video_id: int,
    current_user_role: str = Header("editor")
):
    # CHỈ TƯỜNG VY (ADMIN) MỚI CÓ QUYỀN XÓA VIDEO!
    if current_user_role != "admin":
        raise HTTPException(status_code=403, detail="Chỉ Tường Vy (Quản trị viên) mới có quyền xóa video khỏi hệ thống!")

    filename = database.delete_video(video_id)
    if filename:
        filepath = os.path.join(UPLOAD_DIR, filename)
        if os.path.exists(filepath):
            try:
                os.remove(filepath)
            except Exception:
                pass
    return {"status": "success", "message": "Tường Vy đã xóa video thành công"}

@app.get("/videos/stream/{filename}")
async def stream_video(filename: str):
    safe_filename = os.path.basename(filename)
    filepath = os.path.join(UPLOAD_DIR, safe_filename)
    if not os.path.exists(filepath):
        raise HTTPException(status_code=404, detail="File video không tồn tại")
    
    return FileResponse(
        path=filepath,
        media_type="video/mp4",
        filename=safe_filename
    )

@app.get("/api/export")
async def export_excel(
    tab: str = Query("all"),
    search: Optional[str] = None,
    creator: Optional[str] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None
):
    videos = database.get_all_videos(
        tab=tab,
        search=search,
        creator=creator,
        date_from=date_from,
        date_to=date_to
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Danh_Sach_Video"
    ws.views.sheetView[0].showGridLines = True

    NAVY = "1E3A8A"
    WHITE = "FFFFFF"
    GRAY_LIGHT = "F8FAFC"
    BORDER_COLOR = "E2E8F0"

    font_title = Font(name="Segoe UI", size=15, bold=True, color="1E3A8A")
    font_sub = Font(name="Segoe UI", size=10, italic=True, color="64748B")
    font_header = Font(name="Segoe UI", size=11, bold=True, color=WHITE)
    font_cell = Font(name="Segoe UI", size=10)
    font_cell_bold = Font(name="Segoe UI", size=10, bold=True)

    fill_header = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
    fill_even = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    fill_odd = PatternFill(start_color=GRAY_LIGHT, end_color=GRAY_LIGHT, fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color=BORDER_COLOR),
        right=Side(style='thin', color=BORDER_COLOR),
        top=Side(style='thin', color=BORDER_COLOR),
        bottom=Side(style='thin', color=BORDER_COLOR)
    )

    tab_name_map = {
        "archive": "KHO LƯU TRỮ (CHƯA ĐĂNG)",
        "tiktok": "KÊNH TIKTOK",
        "facebook": "KÊNH FACEBOOK",
        "youtube": "KÊNH YOUTUBE",
        "all": "TẤT CẢ VIDEO"
    }

    ws.merge_cells("A1:H1")
    ws["A1"] = f"BÁO CÁO DỮ LIỆU VIDEO - {tab_name_map.get(tab, 'TỔNG HỢP')}"
    ws["A1"].font = font_title
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    export_time_str = datetime.now().strftime("%d/%m/%Y %H:%M")
    ws.merge_cells("A2:H2")
    ws["A2"] = f"Thời gian xuất: {export_time_str} | Tổng số video: {len(videos)}"
    ws["A2"].font = font_sub
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 18

    headers = [
        "STT", "Ngày Làm", "Tiêu Đề Video", "Người Làm (Editor)",
        "Trạng Thái", "Kênh Đã Đăng", "Ngày Đăng", "File Video"
    ]
    
    header_row = 4
    ws.row_dimensions[header_row].height = 26

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    current_row = 5
    for idx, v in enumerate(videos, start=1):
        fill_current = fill_odd if idx % 2 == 1 else fill_even
        
        is_published = v["status"] == "da_dang" or v["posted_tiktok"] or v["posted_facebook"] or v["posted_youtube"]
        status_text = "Đã Đăng" if is_published else "Lưu Trữ (Chưa đăng)"
        
        channels_text = v["channels"] if is_published else "—"
        has_video = "Đã upload" if v["video_filename"] else "Chưa có"

        row_data = [
            idx,
            v["created_date"] or "",
            v["title"] or "",
            v["creator"] or "",
            status_text,
            channels_text,
            v["published_date"] or "",
            has_video
        ]

        ws.row_dimensions[current_row].height = 22
        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=current_row, column=col_idx, value=value)
            cell.font = font_cell
            cell.fill = fill_current
            cell.border = thin_border
            
            if col_idx in [1, 2, 5, 6, 7, 8]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

            if col_idx == 5:
                cell.font = font_cell_bold
                if is_published:
                    cell.fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
                    cell.font = Font(name="Segoe UI", size=10, bold=True, color="166534")
                else:
                    cell.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
                    cell.font = Font(name="Segoe UI", size=10, bold=True, color="92400E")

        current_row += 1

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row in [1, 2]:
                continue
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    ws.column_dimensions['C'].width = 38
    ws.column_dimensions['F'].width = 24

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename_export = f"Bao_Cao_Video_{tab}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename_export}"}
    )

if __name__ == "__main__":
    if sys.platform.startswith('win'):
        try:
            sys.stdout.reconfigure(encoding='utf-8', errors='replace')
            sys.stderr.reconfigure(encoding='utf-8', errors='replace')
        except Exception:
            pass

    import uvicorn
    port = int(os.environ.get("PORT", 8000))
    ip = get_local_ip()
    print("=" * 65)
    print("   HE THONG KHO & QUAN LY VIDEO - TUONG VY ADMIN")
    print("=" * 65)
    print(f" [+] May Tuong Vy (Chu may):      http://localhost:{port}")
    print(f" [+] Sep & Nhan vien (Wi-Fi LAN): http://{ip}:{port}")
    print("=" * 65)
    print(" [*] Dang tu dong mo trinh duyet web...")
    print(" [!] Giu cua so nay mo trong khi lam viec. Bam [X] de tat.")
    print("=" * 65)
    import webbrowser, threading, time
    def open_browser():
        time.sleep(1.2)
        webbrowser.open(f"http://localhost:{port}")
    threading.Thread(target=open_browser, daemon=True).start()
    uvicorn.run(app, host="0.0.0.0", port=port, reload=False)
